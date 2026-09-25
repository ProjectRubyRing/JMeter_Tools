# -*- coding: utf-8 -*-
"""
解説資料の共通レンダラ（Excel: Meiryo UI・モノトーン / Markdown）。

シート定義:
    {"name": シート名, "title": 見出し, "widths": [列幅...], "blocks": [ブロック...]}
ブロック:
    ("h1", "見出し") / ("h2", "見出し") / ("h3", "見出し")
    ("p", "本文") / ("note", "補足")
    ("art", "図（等幅）")
    ("table", [ヘッダ...], [[セル...], ...])
    ("kv", [[項目, 内容], ...])
    ("items", [列名...], [行...])   行は dict（"group" 行は区切り見出し）
"""
import os
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

FONT = "Meiryo UI"

# ---- モノトーン配色 ---------------------------------------------------------
C_TITLE = "1A1A1A"
C_H1 = "333333"
C_H2 = "595959"
C_H3 = "7F7F7F"
C_HEAD = "404040"
C_GROUP = "D9D9D9"
C_ZEBRA = "F2F2F2"
C_WHITE = "FFFFFF"
C_NOTE = "EDEDED"
C_ART = "F7F7F7"
C_LINE = "BFBFBF"
C_LINE_D = "7F7F7F"
C_TEXT = "000000"
C_SUB = "262626"

THIN = Side(style="thin", color=C_LINE)
THIN_D = Side(style="thin", color=C_LINE_D)
B_CELL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
B_HEAD = Border(left=THIN_D, right=THIN_D, top=THIN_D, bottom=THIN_D)

BASE_SIZE = 10
MAX_ROW_H = 409.0


def disp_len(s):
    """表示幅（全角=2・半角=1）"""
    n = 0
    for ch in s:
        n += 2 if unicodedata.east_asian_width(ch) in ("W", "F", "A") else 1
    return n


# 行高の見積り（Excel 16 の「行の高さの自動調整」で 960 セルを実測して校正）
#   ・Meiryo UI 10pt の 1 行 = 14.4pt
#   ・英数字の連続（URL・ARN・識別子）は単語ごとに次の行へ送られるため、単語単位で折り返しを再現する
#   ・列幅 1 あたり 1 行に入る表示幅（全角=2）は、このモデルで全列 1.25 以上だったため、余裕を見て 1.20 を使う
CHARS_PER_WIDTH = 1.20
LINE_PT_PER_SIZE = 1.44
_TOKEN = re.compile(r"[\x21-\x7e]+|\s|.", re.S)
# 日本語の禁則処理（行頭に置けない文字は前と、行末に置けない文字は後ろと一緒に次の行へ送られる）
_NO_START = set("）」』】〕〉》］｝、。，．・：；？！ー々ゝゞヽヾァィゥェォッャュョヮヵヶぁぃぅぇぉっゃゅょゎ")
_NO_END = set("（「『【〔〈《［｛")


def _tokens(s):
    out = []
    for t in _TOKEN.findall(s):
        if out and not t.isspace() and not out[-1].isspace() and (t[0] in _NO_START or out[-1][-1] in _NO_END):
            out[-1] += t
        else:
            out.append(t)
    return out


def _wrap_count(s, cap):
    lines = 1
    cur = 0.0
    for t in _tokens(s):
        w = disp_len(t)
        if t.isspace():
            if cur + w <= cap:
                cur += w
            continue
        if cur + w <= cap:
            cur += w
        elif w > cap:
            if cur > 0:
                lines += 1
            rem = w
            while rem > cap:
                lines += 1
                rem -= cap
            cur = rem
        else:
            lines += 1
            cur = w
    return lines


def lines_needed(text, width_units, size):
    if text is None or text == "":
        return 1
    cap = max(4.0, (width_units - 1.0) * CHARS_PER_WIDTH * (BASE_SIZE / size))
    return max(1, sum(_wrap_count(raw, cap) for raw in str(text).split("\n")))


def est_height(text, width_units, size=BASE_SIZE, bold=False):
    line_h = size * LINE_PT_PER_SIZE + 0.2
    eff_w = width_units * (0.92 if bold else 1.0)     # 太字は字幅が広いぶん 1 行に入る量を減らす
    return min(MAX_ROW_H, lines_needed(text, eff_w, size) * line_h + 5.0)


def style(c, size=BASE_SIZE, bold=False, color=C_TEXT, bg=None, halign="left", valign="top", wrap=True,
          border=B_CELL, underline=None):
    c.font = Font(name=FONT, size=size, bold=bold, color=color, underline=underline)
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal=halign, vertical=valign, wrap_text=wrap)
    if border is not None:
        c.border = border


def span(ws, row, ncols, text, size, bold, color, bg, border=B_CELL, total_w=100, halign="left", min_h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    style(c, size=size, bold=bold, color=color, bg=bg, border=border, halign=halign, valign="center" if min_h else "top")
    for col in range(2, ncols + 1):
        style(ws.cell(row=row, column=col), size=size, bg=bg, border=border)
    h = est_height(text, total_w, size)
    if min_h:
        h = max(h, min_h)
    ws.row_dimensions[row].height = h
    return row + 1


def _layout(nh, ncols, spans):
    """表の各列が占めるシートの列範囲 [(開始, 終了), ...] を返す（1 始まり）"""
    if spans is None:
        spans = [1] * nh
        spans[-1] = ncols - (nh - 1)
    out = []
    col = 1
    for i, s in enumerate(spans):
        end = ncols if i == nh - 1 else col + s - 1
        out.append((col, end))
        col = end + 1
    return out


def write_table(ws, row, headers, rows, widths, first_bold=True, spans=None):
    ncols = len(widths)
    nh = len(headers)
    lay = _layout(nh, ncols, spans)

    def put(r, i, val, **kw):
        s, e = lay[i]
        if e > s:
            ws.merge_cells(start_row=r, start_column=s, end_row=r, end_column=e)
        style(ws.cell(row=r, column=s, value=val), **kw)
        for col in range(s + 1, e + 1):
            style(ws.cell(row=r, column=col), bg=kw.get("bg"), border=kw.get("border", B_CELL))
        return sum(widths[s - 1:e])

    hh = 24.0
    for i, h in enumerate(headers):
        wu = put(row, i, h, bold=True, color=C_WHITE, bg=C_HEAD, halign="center", valign="center", border=B_HEAD)
        hh = max(hh, est_height(h, wu, bold=True))
    ws.row_dimensions[row].height = hh
    head = row
    row += 1
    for ri, r in enumerate(rows):
        bg = C_WHITE if ri % 2 == 0 else C_ZEBRA
        hmax = 20.0
        for i in range(nh):
            val = r[i] if i < len(r) else ""
            bold = first_bold and i == 0 and nh > 1
            wu = put(row, i, val, bg=bg, bold=bold)
            hmax = max(hmax, est_height(val, wu, bold=bold))
        ws.row_dimensions[row].height = hmax
        row += 1
    return head, row


def build_xlsx(sheets, path, book_title, subject=""):
    wb = Workbook()
    wb.remove(wb.active)
    f0 = Font(name=FONT, size=BASE_SIZE)
    try:
        wb._fonts[0] = f0
    except Exception:
        pass
    try:
        wb._named_styles["Normal"].font = f0
    except Exception:
        pass
    wb.properties.title = book_title
    wb.properties.subject = subject
    wb.properties.creator = "openapi2jmx docs builder"

    # ---- 目次
    toc = wb.create_sheet("00_目次")
    toc.sheet_view.showGridLines = False
    toc.sheet_view.zoomScale = 100
    toc.sheet_properties.tabColor = C_TITLE
    widths = [6, 26, 100]
    for i, w in enumerate(widths, start=1):
        toc.column_dimensions[get_column_letter(i)].width = w
    span(toc, 1, 3, book_title, 14, True, C_WHITE, C_TITLE, B_HEAD, sum(widths), min_h=34)
    if subject:
        span(toc, 2, 3, subject, BASE_SIZE, False, C_SUB, C_NOTE, B_CELL, sum(widths))
    r = 4
    for i, h in enumerate(["No.", "シート名", "内容"], start=1):
        style(toc.cell(row=r, column=i, value=h), bold=True, color=C_WHITE, bg=C_HEAD, halign="center",
              valign="center", border=B_HEAD)
    toc.row_dimensions[r].height = 24
    r += 1
    for i, sh in enumerate(sheets, start=1):
        bg = C_WHITE if i % 2 else C_ZEBRA
        style(toc.cell(row=r, column=1, value=i), bg=bg, halign="center")
        lc = toc.cell(row=r, column=2, value=sh["name"])
        lc.hyperlink = "#'{}'!A1".format(sh["name"])
        style(lc, bold=True, bg=bg, underline="single")
        style(toc.cell(row=r, column=3, value=sh["title"]), bg=bg)
        toc.row_dimensions[r].height = max(22, est_height(sh["title"], 100))
        r += 1
    toc.freeze_panes = "A5"
    _page(toc, landscape=True, paper=9)

    # ---- 本文
    for sh in sheets:
        widths = sh["widths"]
        ncols = len(widths)
        total = sum(widths)
        ws = wb.create_sheet(sh["name"])
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = sh.get("zoom", 100)
        ws.sheet_properties.tabColor = C_H2
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        row = span(ws, 1, ncols, sh["title"], 13, True, C_WHITE, C_TITLE, B_HEAD, total, min_h=30)
        freeze = "A2"
        for blk in sh["blocks"]:
            kind = blk[0]
            if kind in ("h1", "h2", "h3"):
                ws.row_dimensions[row].height = 8
                row += 1
                bg = {"h1": C_H1, "h2": C_H2, "h3": C_H3}[kind]
                size = {"h1": 12, "h2": 11, "h3": 10.5}[kind]
                row = span(ws, row, ncols, blk[1], size, True, C_WHITE, bg, B_HEAD, total, min_h=22)
            elif kind == "p":
                row = span(ws, row, ncols, blk[1], BASE_SIZE, False, C_TEXT, C_WHITE, B_CELL, total)
            elif kind == "note":
                row = span(ws, row, ncols, "【補足】" + blk[1], BASE_SIZE, False, C_SUB, C_NOTE, B_CELL, total)
            elif kind == "art":
                for ln in blk[1].rstrip("\n").split("\n"):
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
                    c = ws.cell(row=row, column=1, value=ln)
                    style(c, size=9.5, color=C_SUB, bg=C_ART, wrap=False, border=None)
                    c.font = Font(name="MS Gothic", size=9.5, color=C_SUB)
                    for col in range(2, ncols + 1):
                        style(ws.cell(row=row, column=col), size=9.5, bg=C_ART, wrap=False, border=None)
                    ws.row_dimensions[row].height = 15
                    row += 1
            elif kind == "kv":
                head, row = write_table(ws, row, ["項目", "内容"], blk[1], widths)
            elif kind == "table":
                opts = blk[3] if len(blk) > 3 else {}
                if isinstance(opts, str):
                    opts = {"filter": opts == "filter"}
                head, row = write_table(ws, row, blk[1], blk[2], widths, spans=opts.get("spans"))
                if opts.get("filter") and ws.auto_filter.ref is None:
                    ws.auto_filter.ref = "A{}:{}{}".format(head, get_column_letter(ncols), row - 1)
                    freeze = "A{}".format(head + 1)
            elif kind == "items":
                cols = blk[1]
                rows = blk[2]
                for i, h in enumerate(cols, start=1):
                    style(ws.cell(row=row, column=i, value=h), bold=True, color=C_WHITE, bg=C_HEAD,
                          halign="center", valign="center", border=B_HEAD)
                ws.row_dimensions[row].height = 30
                head = row
                row += 1
                zebra = 0
                for it in rows:
                    if "group" in it:
                        row = span(ws, row, ncols, it["group"], 10.5, True, C_TEXT, C_GROUP, B_HEAD, total, min_h=22)
                        zebra = 0
                        continue
                    bg = C_WHITE if zebra % 2 == 0 else C_ZEBRA
                    zebra += 1
                    hmax = 22.0
                    for i, key in enumerate(blk[3], start=1):
                        val = it.get(key, "")
                        c = ws.cell(row=row, column=i, value=val)
                        bold = key in ("no", "path")
                        style(c, bg=bg, bold=bold, halign="center" if key == "no" else "left")
                        hmax = max(hmax, est_height(val, widths[i - 1], bold=bold))
                    ws.row_dimensions[row].height = hmax
                    row += 1
                ws.auto_filter.ref = "A{}:{}{}".format(head, get_column_letter(ncols), row - 1)
                freeze = sh.get("freeze", "D{}".format(head + 1))
        ws.freeze_panes = freeze
        _page(ws, landscape=True, paper=sh.get("paper", 8))
        ws.print_title_rows = "1:1"
    d = os.path.dirname(path)
    if d:
        os.makedirs(d, exist_ok=True)
    wb.save(path)
    return path


def _page(ws, landscape=True, paper=9):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = paper       # 8=A3, 9=A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "&P / &N"
    ws.oddFooter.center.font = FONT


# ============================================================================
# Markdown
# ============================================================================
def md_cell(s):
    return str(s).replace("|", "\\|").replace("\n", "<br>")


def anchor_of(text):
    a = re.sub(r"[^\w\-ぁ-んァ-ヶー一-龠]", "", text.replace(" ", "-"))
    return a.lower()


def build_md(sheets, path, book_title, lead=""):
    out = ["# " + book_title, ""]
    if lead:
        out += [lead, ""]
    out += ["## 目次", ""]
    for i, sh in enumerate(sheets, start=1):
        out.append("{}. [{}](#{}) — {}".format(i, sh["name"], "sheet-{}".format(i), sh["title"]))
    out += ["", "---", ""]
    for i, sh in enumerate(sheets, start=1):
        out += ['<a id="sheet-{}"></a>'.format(i), "", "# {}　{}".format(sh["name"], sh["title"]), ""]
        for blk in sh["blocks"]:
            kind = blk[0]
            if kind == "h1":
                out += ["## " + blk[1], ""]
            elif kind == "h2":
                out += ["### " + blk[1], ""]
            elif kind == "h3":
                out += ["#### " + blk[1], ""]
            elif kind == "p":
                out += [blk[1], ""]
            elif kind == "note":
                out += ["> " + ln for ln in ("【補足】" + blk[1]).split("\n")] + [""]
            elif kind == "art":
                out += ["```text", blk[1].rstrip("\n"), "```", ""]
            elif kind in ("table", "kv"):
                headers, rows = (["項目", "内容"], blk[1]) if kind == "kv" else (blk[1], blk[2])
                out.append("| " + " | ".join(md_cell(h) for h in headers) + " |")
                out.append("|" + "|".join([" --- "] * len(headers)) + "|")
                for r in rows:
                    out.append("| " + " | ".join(md_cell(r[j]) if j < len(r) else "" for j in range(len(headers))) + " |")
                out.append("")
            elif kind == "items":
                labels = dict(zip(blk[3], blk[1]))
                for it in blk[2]:
                    if "group" in it:
                        out += ["## " + it["group"], ""]
                        continue
                    out += ["### {}　`{}`".format(it["no"], it["path"]), ""]
                    out.append("| 項目 | 内容 |")
                    out.append("| --- | --- |")
                    for key in blk[3]:
                        if key in ("no", "path"):
                            continue
                        val = it.get(key, "")
                        if key == "value":
                            val = val if val else "（オブジェクト）"
                        out.append("| **{}** | {} |".format(md_cell(labels[key]), md_cell(val)))
                    out.append("")
        out += ["---", ""]
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(out))
    return path
